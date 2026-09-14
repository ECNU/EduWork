"""Actual 4:3 PPTX exercising file properties independently of our generator."""
import sys
from io import BytesIO
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
page = deck.slides.add_slide(deck.slide_layouts[6])
card = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.5), Inches(.5), Inches(4), Inches(2))
card.name = 'Rounded file card'
card.fill.solid(); card.fill.fore_color.rgb = RGBColor(240, 245, 252)
card.line.color.rgb = RGBColor(40, 80, 120); card.line.width = Pt(2)
text = page.shapes.add_textbox(Inches(.75), Inches(.75), Inches(3.5), Inches(1.5))
text.name = 'Explicit font and breaks'
text.text_frame.text = 'Original page\n中文手动换行\v第二行'
for paragraph in text.text_frame.paragraphs:
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = 'Microsoft YaHei'; run.font.size = Pt(20)
single = page.shapes.add_textbox(Inches(5), Inches(.75), Inches(4), Inches(.6))
single.name = 'No wrap value'
single.text_frame.word_wrap = False
single.text = '100%→99%'
single.text_frame.paragraphs[0].runs[0].font.name = 'Georgia'
single.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
circle = page.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.5), Inches(2), Inches(1), Inches(1))
circle.name = 'Circle file shape'
circle.fill.solid(); circle.fill.fore_color.rgb = RGBColor(20, 140, 100)
circle.line.fill.background()
picture = Image.new('RGB', (200, 100), 'red')
for x in range(100, 200):
    for y in range(100):
        picture.putpixel((x, y), (0, 80, 200))
blob = BytesIO(); picture.save(blob, format='PNG'); blob.seek(0)
image = page.shapes.add_picture(blob, Inches(1), Inches(3.5), width=Inches(2), height=Inches(2))
image.name = 'Cropped image'; image.crop_left = .5
table = page.shapes.add_table(2, 2, Inches(4), Inches(4), Inches(5), Inches(1.5)).table
table.columns[0].width = Inches(3); table.columns[1].width = Inches(2)
for cell, value in zip([c for row in table.rows for c in row.cells], ['资料', '数量', '合成样本', '28']):
    cell.text = value
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Microsoft YaHei'; run.font.size = Pt(16)
page = deck.slides.add_slide(deck.slide_layouts[6])
data = CategoryChartData(); data.categories = ['A', 'B']; data.add_series('合成数据', (2, 4))
page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), data)
output = Path(sys.argv[1]); assert not output.exists()
deck.save(output)
