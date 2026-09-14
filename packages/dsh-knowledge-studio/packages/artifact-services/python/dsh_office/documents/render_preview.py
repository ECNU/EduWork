from __future__ import annotations

import argparse
import base64
import hashlib
import html
import json
import math
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


MAX_SHEETS = 12
MAX_ROWS = 240
MAX_COLUMNS = 60
MAX_SLIDES = 80


def escaped(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def data_url(content_type: str, blob: bytes) -> str:
    return f"data:{content_type};base64,{base64.b64encode(blob).decode('ascii')}"


def color_value(color: Any) -> str | None:
    try:
        value = str(color.rgb or "")
    except (AttributeError, TypeError, ValueError):
        return None
    if len(value) == 8:
        value = value[2:]
    return f"#{value}" if len(value) == 6 else None


def document_shell(title: str, kind: str, body: str, extra_style: str = "") -> str:
    safe_title = escaped(title)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src data:">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{safe_title}</title><style>
:root{{color-scheme:light;--paper:#fff;--ink:#252525;--muted:#6d7480;--line:#d9dde5;--canvas:#eef0f4;--accent:#8b2232}}
*{{box-sizing:border-box}}html,body{{margin:0;min-height:100%;background:var(--canvas);color:var(--ink);font-family:"Microsoft YaHei","Segoe UI",sans-serif}}
body{{padding:18px}}.preview-meta{{max-width:1100px;margin:0 auto 12px;color:var(--muted);font-size:12px}}
.paper{{width:min(100%,860px);min-height:1080px;margin:0 auto 18px;padding:58px 68px;background:var(--paper);box-shadow:0 5px 24px #1f29371f}}
.paper p{{margin:.45em 0;line-height:1.72;white-space:pre-wrap}}.paper h1,.paper h2,.paper h3{{line-height:1.35;margin:1em 0 .5em}}
.paper img{{max-width:100%;height:auto}}table{{border-collapse:collapse}}.paper table{{width:100%;margin:14px 0}}th,td{{border:1px solid var(--line);padding:6px 8px;vertical-align:top}}
.page-break{{border:0;border-top:2px dashed var(--line);margin:34px -20px}}.note{{padding:9px 12px;border-radius:8px;background:#fff8e7;color:#705a1d;font-size:12px}}
.sheet{{width:max-content;min-width:min(100%,900px);max-width:none;margin:0 auto 22px;background:#fff;box-shadow:0 4px 18px #1f29371a}}
.sheet h2{{position:sticky;left:0;margin:0;padding:12px 14px;border-bottom:1px solid var(--line);font-size:15px}}.sheet-wrap{{max-width:100%;overflow:auto}}
.sheet table{{font:12px/1.45 "Segoe UI",sans-serif}}.sheet td{{min-width:64px;max-width:360px;white-space:pre-wrap;overflow-wrap:anywhere}}
.row-head{{position:sticky;left:0;z-index:2;min-width:42px!important;background:#f5f6f8!important;color:var(--muted);text-align:right}}
.slides{{display:grid;gap:24px;justify-items:center}}.slide-wrap{{width:min(100%,1040px)}}.slide-label{{margin:0 0 7px;color:var(--muted);font-size:12px}}
.slide{{position:relative;container-type:inline-size;width:100%;overflow:hidden;background:#fff;box-shadow:0 5px 24px #1f29372b}}.shape{{position:absolute;overflow:hidden}}
.shape-text{{display:flex;flex-direction:column;overflow-wrap:anywhere;line-height:1.12}}.shape-text p{{width:100%;margin:0;white-space:pre-wrap}}.shape img{{display:block;width:100%;height:100%;object-fit:contain}}
.shape table{{width:100%;height:100%;table-layout:fixed;font-size:1.04cqw;background:#fff}}.shape td{{padding:.2cqw .4cqw;overflow:hidden}}
{extra_style}
</style></head><body><div class="preview-meta">{escaped(kind)} · 隔离预览，不执行文档中的脚本或外部资源</div>{body}</body></html>"""


def run_html(run: Any) -> str:
    styles: list[str] = []
    if run.bold:
        styles.append("font-weight:700")
    if run.italic:
        styles.append("font-style:italic")
    if run.underline:
        styles.append("text-decoration:underline")
    if run.font.size is not None:
        styles.append(f"font-size:{run.font.size.pt:.2f}pt")
    color = color_value(run.font.color)
    if color:
        styles.append(f"color:{color}")
    parts = [escaped(run.text).replace("\n", "<br>")]
    try:
        relationships = run.part.related_parts
        for blip in run._r.xpath(".//a:blip"):
            relation_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = relationships.get(relation_id)
            if part is not None and hasattr(part, "blob"):
                content_type = getattr(part, "content_type", "image/png")
                parts.append(f'<img src="{data_url(content_type, part.blob)}" alt="文档图片">')
    except (AttributeError, KeyError, TypeError, ValueError):
        pass
    return f'<span style="{";".join(styles)}">{"".join(parts)}</span>'


def paragraph_html(paragraph: Any) -> str:
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    tag = "p"
    if "heading 1" in style_name or "标题 1" in style_name:
        tag = "h1"
    elif "heading 2" in style_name or "标题 2" in style_name:
        tag = "h2"
    elif "heading 3" in style_name or "标题 3" in style_name:
        tag = "h3"
    content = "".join(run_html(run) for run in paragraph.runs) or escaped(paragraph.text)
    extra = ' class="page-break"' if "page break" in style_name else ""
    if extra:
        return f"<hr{extra}>"
    return f"<{tag}>{content}</{tag}>"


def word_table_html(table: Any) -> str:
    rows = []
    for row in table.rows:
        cells = "".join(f"<td>{''.join(paragraph_html(p) for p in cell.paragraphs)}</td>" for cell in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return f"<table><tbody>{''.join(rows)}</tbody></table>"


def render_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            parts.append(paragraph_html(Paragraph(child, document)))
        elif child.tag.endswith("}tbl"):
            parts.append(word_table_html(Table(child, document)))
    return document_shell(path.name, "Word 文档", f'<main class="paper">{"".join(parts)}</main>')


def cell_css(cell: Any) -> str:
    styles: list[str] = []
    fill = color_value(cell.fill.fgColor)
    font = color_value(cell.font.color)
    if fill and fill.lower() not in {"#000000", "#ffffff"}:
        styles.append(f"background:{fill}")
    if font:
        styles.append(f"color:{font}")
    if cell.font.bold:
        styles.append("font-weight:700")
    if cell.font.italic:
        styles.append("font-style:italic")
    if cell.alignment.horizontal in {"center", "right"}:
        styles.append(f"text-align:{cell.alignment.horizontal}")
    if cell.alignment.vertical in {"center", "top", "bottom"}:
        styles.append(f"vertical-align:{'middle' if cell.alignment.vertical == 'center' else cell.alignment.vertical}")
    return ";".join(styles)


def cell_text(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return escaped(value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat())
    return escaped(value).replace("\n", "<br>")


def render_xlsx(path: Path) -> str:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(str(path), read_only=False, data_only=False)
    sections: list[str] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets[:MAX_SHEETS], start=1):
            max_row = min(max(worksheet.max_row, 1), MAX_ROWS)
            max_column = min(max(worksheet.max_column, 1), MAX_COLUMNS)
            merged_starts: dict[tuple[int, int], tuple[int, int]] = {}
            merged_covered: set[tuple[int, int]] = set()
            for area in worksheet.merged_cells.ranges:
                if area.min_row > max_row or area.min_col > max_column:
                    continue
                row_span = min(area.max_row, max_row) - area.min_row + 1
                col_span = min(area.max_col, max_column) - area.min_col + 1
                merged_starts[(area.min_row, area.min_col)] = (row_span, col_span)
                for row in range(area.min_row, min(area.max_row, max_row) + 1):
                    for column in range(area.min_col, min(area.max_col, max_column) + 1):
                        if (row, column) != (area.min_row, area.min_col):
                            merged_covered.add((row, column))
            colgroup = []
            for column in range(1, max_column + 1):
                width = worksheet.column_dimensions[get_column_letter(column)].width or 10
                colgroup.append(f'<col style="width:{min(max(width * 7, 54), 280):.0f}px">')
            rows: list[str] = []
            for row_index in range(1, max_row + 1):
                cells = [f'<th class="row-head">{row_index}</th>']
                for column_index in range(1, max_column + 1):
                    if (row_index, column_index) in merged_covered:
                        continue
                    cell = worksheet.cell(row_index, column_index)
                    if isinstance(cell, MergedCell):
                        continue
                    span = merged_starts.get((row_index, column_index), (1, 1))
                    attributes = []
                    if span[0] > 1:
                        attributes.append(f'rowspan="{span[0]}"')
                    if span[1] > 1:
                        attributes.append(f'colspan="{span[1]}"')
                    css = cell_css(cell)
                    if css:
                        attributes.append(f'style="{css}"')
                    cells.append(f'<td {" ".join(attributes)}>{cell_text(cell.value)}</td>')
                rows.append(f"<tr>{''.join(cells)}</tr>")
            truncated = worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS
            note = '<p class="note">工作表较大，预览仅显示前 240 行 × 60 列；完整内容请使用“打开”。</p>' if truncated else ""
            state = "" if worksheet.sheet_state == "visible" else f" · {escaped(worksheet.sheet_state)}"
            sections.append(f'<section class="sheet"><h2>{sheet_index}. {escaped(worksheet.title)}{state}</h2>{note}<div class="sheet-wrap"><table><colgroup><col style="width:42px">{"".join(colgroup)}</colgroup><tbody>{"".join(rows)}</tbody></table></div></section>')
    finally:
        workbook.close()
    if len(workbook.sheetnames) > MAX_SHEETS:
        sections.append('<p class="note">工作簿包含较多工作表，预览仅显示前 12 个。</p>')
    return document_shell(path.name, "Excel 工作簿", f'<main>{"".join(sections)}</main>')


def ppt_fill(shape: Any) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        return color_value(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def ppt_length(emu: float, slide_width: int) -> str:
    """Fixed 96-dpi page coordinates; the viewer scales the entire page."""
    return f"{float(emu) / 9525:.6f}px"


def ppt_font(points: float, slide_width: int) -> str:
    return ppt_length(float(points) * 12700, slide_width)


def ppt_font_style(font: Any, slide_width: int) -> list[str]:
    styles = []
    if font.name:
        styles.append(f"font-family:{json.dumps(font.name, ensure_ascii=False)},sans-serif")
    if font.size is not None:
        styles.append(f"font-size:{ppt_font(font.size.pt, slide_width)}")
    if font.bold is not None:
        styles.append("font-weight:700" if font.bold else "font-weight:400")
    if font.italic is not None:
        styles.append("font-style:italic" if font.italic else "font-style:normal")
    if font.underline:
        styles.append("text-decoration:underline")
    color = color_value(font.color)
    if color:
        styles.append(f"color:{color}")
    return styles


def ppt_text(shape: Any, slide_width: int) -> tuple[str, str]:
    paragraphs: list[str] = []
    styles: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        paragraph_styles = []
        alignment = str(paragraph.alignment or "").lower()
        for value in ("left", "center", "right", "justify"):
            if alignment.startswith(value):
                paragraph_styles.append(f"text-align:{value}")
                break
        paragraph_styles.extend(ppt_font_style(paragraph.font, slide_width))
        for field, css in (("space_before", "margin-top"), ("space_after", "margin-bottom")):
            spacing = getattr(paragraph, field, None)
            if spacing is not None:
                paragraph_styles.append(f"{css}:{ppt_length(spacing, slide_width)}")
        spacing = paragraph.line_spacing
        if spacing is not None:
            line_height = ppt_length(spacing, slide_width) if hasattr(spacing, "pt") else str(float(spacing))
            paragraph_styles.append(f"line-height:{line_height}")
        runs = []
        run_iterator = iter(paragraph.runs)
        # Iterating only .runs loses OOXML soft breaks between runs.
        for child in paragraph._p:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "br":
                runs.append("<br>")
            elif tag == "r":
                run = next(run_iterator)
                run_styles = ppt_font_style(run.font, slide_width)
                runs.append(f'<span style="{escaped(";".join(run_styles))}">{escaped(run.text)}</span>')
            elif tag == "fld":
                runs.append(escaped("".join(child.itertext())))
        paragraphs.append(f'<p style="{escaped(";".join(paragraph_styles))}">{"".join(runs) or escaped(paragraph.text).replace(chr(11), "<br>")}</p>')
    first_run = next((run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text), None)
    if first_run is not None and first_run.font.size is not None:
        styles.append(f"font-size:{ppt_font(first_run.font.size.pt, slide_width)}")
    else:
        styles.append(f"font-size:{ppt_font(18, slide_width)}")
    for side in ("left", "right", "top", "bottom"):
        margin = getattr(shape.text_frame, f"margin_{side}", 0) or 0
        styles.append(f"padding-{side}:{ppt_length(margin, slide_width)}")
    anchor = str(getattr(shape.text_frame, "vertical_anchor", "")).lower()
    styles.append("justify-content:center" if "middle" in anchor else "justify-content:flex-end" if "bottom" in anchor else "justify-content:flex-start")
    if shape.text_frame.word_wrap is False:
        styles.append("--ppt-white-space:pre;overflow-wrap:normal")
    return "".join(paragraphs), ";".join(styles)


def ppt_table(shape: Any, slide_width: int) -> str:
    rows = []
    for row in shape.table.rows:
        cells = []
        for cell in row.cells:
            if cell.is_spanned:
                continue
            content, style = ppt_text(cell, slide_width)
            fill = ppt_fill(cell)
            for side in ("left", "right", "top", "bottom"):
                style += f";padding-{side}:{ppt_length(getattr(cell, f'margin_{side}', 0) or 0, slide_width)}"
            cells.append(f'<td rowspan="{cell.span_height}" colspan="{cell.span_width}" style="background:{fill or "transparent"}"><div class="shape-text" style="{escaped(style)}">{content}</div></td>')
        rows.append(f'<tr style="height:{ppt_length(row.height, slide_width)}">{"".join(cells)}</tr>')
    columns = "".join(f'<col style="width:{ppt_length(column.width, slide_width)}">' for column in shape.table.columns)
    return f"<table><colgroup>{columns}</colgroup><tbody>{''.join(rows)}</tbody></table>"


def ppt_nodes(node: Any, expression: str) -> list[Any]:
    from lxml.etree import XPath
    return XPath(expression, namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"})(node)


def ppt_shadow(shape: Any, theme: Any) -> str | None:
    shadows = ppt_nodes(shape._element, "./p:spPr/a:effectLst/a:outerShdw")
    if not shadows and not ppt_nodes(shape._element, "./p:spPr/a:effectLst") and theme is not None:
        refs = ppt_nodes(shape._element, "./p:style/a:effectRef")
        if refs:
            index = int(refs[0].get("idx", "0")) - 1
            effects = ppt_nodes(theme, ".//a:effectStyleLst/a:effectStyle")
            if 0 <= index < len(effects):
                shadows = ppt_nodes(effects[index], "./a:effectLst/a:outerShdw")
    if not shadows:
        return None
    shadow = shadows[0]
    colors = ppt_nodes(shadow, "./a:srgbClr")
    if not colors:
        return None
    color = colors[0].get("val", "")
    if len(color) != 6 or any(c not in "0123456789abcdefABCDEF" for c in color):
        return None
    alphas = ppt_nodes(colors[0], "./a:alpha")
    alpha = max(0, min(1, float(alphas[0].get("val", "100000")) / 100000)) if alphas else 1
    distance = float(shadow.get("dist", "0")) / 9525
    angle = math.radians(float(shadow.get("dir", "0")) / 60000)
    blur = float(shadow.get("blurRad", "0")) / 9525
    rgb = ",".join(str(int(color[index:index+2], 16)) for index in (0, 2, 4))
    return f"box-shadow:{distance*math.cos(angle):.3f}px {distance*math.sin(angle):.3f}px {blur:.3f}px rgba({rgb},{alpha:.3f})"


def shape_html(shape: Any, slide_width: int, slide_height: int, z_index: int, image_styles: dict[str, str], warnings: set[tuple[str, str]], theme: Any) -> str:
    rotation = float(getattr(shape, "rotation", 0) or 0)
    base = [
        f"left:{ppt_length(shape.left, slide_width)}", f"top:{ppt_length(shape.top, slide_width)}",
        f"width:{ppt_length(shape.width, slide_width)}", f"height:{ppt_length(shape.height, slide_width)}",
        f"z-index:{z_index}", f"transform:rotate({rotation:.2f}deg)",
    ]
    fill = ppt_fill(shape)
    if fill:
        base.append(f"background:{fill}")
    content = ""
    extra_class = ""
    try:
        shadow = ppt_shadow(shape, theme)
        if shadow:
            base.append(shadow)
        geometry = shape._element.xpath("./p:spPr/a:prstGeom")
        preset = geometry[0].get("prst") if geometry else None
        if preset == "ellipse":
            base.append("border-radius:50%")
        elif preset == "roundRect":
            radius = 1 / 6
            adjustments = geometry[0].xpath("./a:avLst/a:gd[@name='adj']")
            if adjustments:
                formula = adjustments[0].get("fmla", "").split()
                if len(formula) == 2 and formula[0] == "val":
                    radius = max(0, min(.5, float(formula[1]) / 100000))
            base.append(f"border-radius:{ppt_length(min(shape.width, shape.height)*radius, slide_width)}")
        elif preset not in {None, "rect", "line"}:
            warnings.add(("shape-geometry", "部分特殊形状无法准确还原，请打开原始 PPTX 查看。"))
        try:
            stroke = color_value(shape.line.color)
            if stroke:
                stroke_width = ppt_length(shape.line.width or 12700, slide_width)
                if int(shape.shape_type) == 9:
                    base.extend([f"border-top:{stroke_width} solid {stroke}", "overflow:visible"])
                else:
                    base.append(f"outline:{stroke_width} solid {stroke}")
        except (AttributeError, TypeError, ValueError):
            pass
        if hasattr(shape, "image"):
            blob = shape.image.blob
            class_name = f"ppt-image-{hashlib.sha256(blob).hexdigest()[:16]}"
            image_styles.setdefault(class_name, data_url(shape.image.content_type, blob))
            left, right = shape.crop_left, shape.crop_right
            top, bottom = shape.crop_top, shape.crop_bottom
            visible_width, visible_height = max(.001, 1-left-right), max(.001, 1-top-bottom)
            content = f'<div class="{class_name}" style="position:absolute;left:{-100*left/visible_width:.6f}%;top:{-100*top/visible_height:.6f}%;width:{100/visible_width:.6f}%;height:{100/visible_height:.6f}%"></div>'
        elif getattr(shape, "has_table", False):
            content = ppt_table(shape, slide_width)
            warnings.add(("table-style", "表格内容与单元格尺寸已读取；主题边框等细节可能与原文件不同。"))
        elif getattr(shape, "has_chart", False) or int(shape.shape_type) in {6, 7, 10, 16, 19, 21}:
            warnings.add(("unsupported-object", "文件含当前轻量预览不支持的图表、组合或嵌入对象，请打开原始 PPTX 查看完整内容。"))
            content = '<span class="unsupported-object">此对象请在原始 PPTX 中查看</span>'
            base.append("outline:1px dashed #a76d27")
        elif getattr(shape, "has_text_frame", False) and shape.text_frame.text:
            content, text_style = ppt_text(shape, slide_width)
            base.append(text_style)
            extra_class = " shape-text"
            if not any(run.font.size is not None for paragraph in shape.text_frame.paragraphs for run in paragraph.runs):
                warnings.add(("inherited-text-style", "文件含继承自主题或母版的文字样式，字号等细节可能与原文件不同。"))
    except (AttributeError, TypeError, ValueError):
        content = '<span class="unsupported-object">此对象预览失败，请打开原始 PPTX 查看</span>'
        warnings.add(("object-conversion", "部分对象转换失败，请打开原始 PPTX 查看完整内容。"))
    return f'<div class="shape{extra_class}" data-shape-name="{escaped(shape.name)}" style="{escaped(";".join(filter(None, base)))}">{content}</div>'


def render_pptx(path: Path) -> str:
    from pptx import Presentation
    from pptx.oxml import parse_xml

    presentation = Presentation(str(path))
    ratio = presentation.slide_width / presentation.slide_height
    width_in = presentation.slide_width / 914400
    height_in = presentation.slide_height / 914400
    slides: list[str] = []
    image_styles: dict[str, str] = {}
    warnings: set[tuple[str, str]] = set()
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            break
        theme = None
        for relation in slide.slide_layout.slide_master.part.rels.values():
            if relation.reltype.endswith("/theme") and not relation.is_external:
                theme = parse_xml(relation.target_part.blob)
                break
        shapes = "".join(shape_html(shape, presentation.slide_width, presentation.slide_height, z, image_styles, warnings, theme) for z, shape in enumerate(slide.shapes, start=1))
        background = ""
        try:
            color = color_value(slide.background.fill.fore_color)
            if color:
                background = f";background:{color}"
        except (AttributeError, TypeError, ValueError):
            pass
        slides.append(f'<section class="slide-wrap" data-slide-index="{index}"><p class="slide-label">第 {index} 页</p><div class="slide" style="width:{ppt_length(presentation.slide_width, presentation.slide_width)};height:{ppt_length(presentation.slide_height, presentation.slide_width)};aspect-ratio:{ratio:.8f}{background}">{shapes}</div></section>')
    if len(presentation.slides) > MAX_SLIDES:
        warnings.add(("page-limit", "演示文稿较长，预览仅显示前 80 页；完整内容请打开原始 PPTX。"))
    picture_css = "".join(f'.{name}{{background-image:url("{source}");background-size:100% 100%;background-repeat:no-repeat;background-position:center}}' for name, source in image_styles.items())
    print_css = f"""
.slides{{display:grid;justify-items:start}}.slide-wrap{{width:{width_in*96:.6f}px}}
.slide{{container-type:normal}}.shape-text{{overflow-wrap:normal;word-break:normal}}
.shape-text p{{white-space:var(--ppt-white-space,pre-wrap);flex:none}}
.shape table{{font-size:13.333333px;background:transparent}}.shape td{{padding:0}}
.unsupported-object{{display:block;padding:8px;font-size:14px;color:#815919}}
@page{{size:{width_in:.6f}in {height_in:.6f}in;margin:0}}
@media print{{
body{{padding:0;background:white}}.preview-meta,.slide-label,.note{{display:none}}
.slides{{display:block}}.slide-wrap{{width:{width_in:.6f}in;height:{height_in:.6f}in;margin:0;break-after:page;break-inside:avoid}}
.slide-wrap:last-child{{break-after:auto}}.slide{{width:100%;height:100%;aspect-ratio:auto!important;box-shadow:none}}
}}
"""
    attributes = f'data-slide-width="{presentation.slide_width / 9525:.6f}" data-slide-height="{presentation.slide_height / 9525:.6f}" data-slide-count="{len(presentation.slides)}"'
    notices = "".join(f'<p class="note" data-office-warning data-office-warning-code="{code}">{escaped(message)}</p>' for code, message in sorted(warnings))
    return document_shell(path.name, "PowerPoint 演示文稿", f'{notices}<main class="slides" {attributes}>{"".join(slides)}</main>', picture_css + print_css)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Render a bounded, static Office preview as HTML.")
    parser.add_argument("--input", required=True)
    return parser.parse_args()


def main() -> int:
    path = Path(parse_args().input).resolve(strict=True)
    suffix = path.suffix.lower()
    renderers = {".docx": render_docx, ".xlsx": render_xlsx, ".pptx": render_pptx}
    renderer = renderers.get(suffix)
    if renderer is None:
        raise ValueError(f"unsupported Office preview type: {suffix}")
    sys.stdout.write(renderer(path))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        raise SystemExit(2)
